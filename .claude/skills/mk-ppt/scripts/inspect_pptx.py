"""PPTX 구조와 텍스트를 출력한다.

사용법: python inspect_pptx.py <파일.pptx>
"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu


def walk(shapes, depth=1):
    pad = "  " * depth
    for sh in shapes:
        kind = sh.shape_type
        print(f"{pad}- [{kind}] {sh.name}")
        if kind == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, depth + 1)
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            for line in sh.text_frame.text.splitlines():
                print(f"{pad}    | {line}")
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                print(f"{pad}    | " + " | ".join(c.text for c in row.cells))
        if getattr(sh, "has_chart", False) and sh.has_chart:
            plot = sh.chart.plots[0]
            print(f"{pad}    차트: {sh.chart.chart_type}, 항목 {list(plot.categories)}")
            for s in plot.series:
                print(f"{pad}    계열 {s.name}: {list(s.values)}")


def main(path):
    prs = Presentation(path)
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    print(f"파일: {path}")
    print(f"슬라이드 크기: {w:.2f} x {h:.2f} in, 슬라이드 수: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n## 슬라이드 {i} (레이아웃: {slide.slide_layout.name})")
        walk(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print(f"  노트: {notes}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
