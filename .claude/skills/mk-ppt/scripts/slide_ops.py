"""python-pptx에 공식 API가 없는 슬라이드 삭제·복제·이동. 번호는 1부터.

사용법:
  python slide_ops.py delete    <in.pptx> <out.pptx> <번호>
  python slide_ops.py duplicate <in.pptx> <out.pptx> <번호>
  python slide_ops.py move      <in.pptx> <out.pptx> <원래번호> <새번호>
"""
import copy
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart

R_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"))


def _remap_rids(element, rid_map):
    for node in element.iter():
        for attr in R_ATTRS:
            if node.get(attr) in rid_map:
                node.set(attr, rid_map[node.get(attr)])


def _clone_chart(slide_part, old_chart):
    """차트 파트와 내장 엑셀 데이터를 새로 복사해 슬라이드에 연결하고 새 rId를 반환한다.

    차트를 공유하면 복제본 차트를 고칠 때 원본도 바뀌므로 독립 복사한다.
    새 파트를 즉시 연결해야 다음 파트 이름(chartN.xml)이 겹치지 않는다.
    """
    package = slide_part.package
    partname = package.next_partname(ChartPart.partname_template)
    new_chart = ChartPart.load(partname, old_chart.content_type, package, old_chart.blob)
    new_rid = slide_part.relate_to(new_chart, RT.CHART)
    rid_map = {}
    for rid, rel in old_chart.rels.items():
        if rel.is_external:
            rid_map[rid] = new_chart.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.PACKAGE:
            xlsx = EmbeddedXlsxPart.new(rel.target_part.blob, package)
            rid_map[rid] = new_chart.relate_to(xlsx, RT.PACKAGE)
        else:
            rid_map[rid] = new_chart.relate_to(rel.target_part, rel.reltype)
    _remap_rids(new_chart._element, rid_map)
    return new_rid


def delete(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[idx]
    prs.part.drop_rel(sldId.get(qn("r:id")))
    sldIdLst.remove(sldId)


def duplicate(prs, idx):
    src = prs.slides[idx]
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    # 이미지·차트 등 관계를 먼저 복사하고 rId를 새 값으로 매핑
    # 노트는 별도 파트라 관계를 건너뛰고 아래에서 텍스트만 복사한다
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT.NOTES_SLIDE, RT.SLIDE_LAYOUT):
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.CHART:
            rid_map[rid] = _clone_chart(new.part, rel.target_part)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new_el = copy.deepcopy(el)
        _remap_rids(new_el, rid_map)
        new.shapes._spTree.append(new_el)
    if src.has_notes_slide:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    if src.background.fill.type is not None:
        bg = src._element.cSld.bg
        if bg is not None:
            new._element.cSld.insert(0, copy.deepcopy(bg))
    # 새 슬라이드는 맨 끝에 추가되므로 원본 바로 뒤로 이동
    move(prs, len(prs.slides) - 1, idx + 1)


def move(prs, old, new):
    sldIdLst = prs.slides._sldIdLst
    el = sldIdLst[old]
    sldIdLst.remove(el)
    sldIdLst.insert(new, el)


def main(argv):
    if len(argv) < 5:
        sys.exit(__doc__)
    op, src, dst = argv[1], argv[2], argv[3]
    nums = [int(x) - 1 for x in argv[4:]]
    prs = Presentation(src)
    n = len(prs.slides)
    if any(not 0 <= x < n for x in nums):
        sys.exit(f"슬라이드 번호는 1~{n} 범위여야 합니다.")
    if op == "delete":
        delete(prs, nums[0])
    elif op == "duplicate":
        duplicate(prs, nums[0])
    elif op == "move" and len(nums) == 2:
        move(prs, nums[0], nums[1])
    else:
        sys.exit(__doc__)
    prs.save(dst)
    print(f"{op} 완료: {dst} (슬라이드 {len(prs.slides)}장)")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv)
